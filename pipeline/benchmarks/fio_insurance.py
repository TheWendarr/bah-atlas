"""
fio_insurance.py — read the A8 Treasury FIO workbook (Supporting Underlying Metrics,
Analyses of U.S. Homeowners Insurance Markets, 2018-2022).

Layout (checked 18 Oct 2026 against the file Treasury published in Jan 2025):
  sheet 'Supporting Underlying Metrics', header in row 1:
    ZIP Code | Year | Policy Decile Grouping | Claim Frequency | Claim Severity |
    Loss Ratio | Premiums Per Policy | Nonrenewal Rate | Nonpayment Cancellation Rate |
    Other than Nonpayment Cancellation Rate
  one row per ZIP and year (2018-2022). ZIPs are stored as numbers, so leading
  zeros are lost (1001 is 01001). Blank rows follow the data. The other two sheets
  ('Disclaimer & Background', 'Metric Definitions') are documentation.

Premiums Per Policy = written premium / policies in force at the end of the year
(FIO's definition). Only ZIPs with at least 10 reporting insurers and 50 policies
are published. FIO notes the ZIP-level values carry anomalies it did not validate
one by one, which is one reason the MHA value is a median over ZIPs.

openpyxl only (already in requirements for B10).
"""

from __future__ import annotations

import statistics
import sys
from pathlib import Path

import mha_agg as agg

SHEET = "Supporting Underlying Metrics"
ZIP_COL, YEAR_COL, PREMIUM_COL = "ZIP Code", "Year", "Premiums Per Policy"
HIGH_PREMIUM = 10_000   # reported, not removed: real in some coastal and high-value ZIPs


def read(path: Path, year: int, tool: str) -> tuple[dict[str, float], dict]:
    """{zip5: premiums per policy} for `year`, plus stats for the manifest."""
    try:
        import openpyxl
    except ImportError:
        agg.die(tool, "openpyxl is required to read the FIO workbook (pip install -r requirements.txt)")
    if not path.exists():
        agg.die(tool, f"FIO workbook not found: {path} (run `make fetch`, or save it there from a browser)")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        agg.die(tool, f"{path.name} has no sheet {SHEET!r}; sheets are {wb.sheetnames}")
    rows = wb[SHEET].iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows, ())]
    missing = [c for c in (ZIP_COL, YEAR_COL, PREMIUM_COL) if c not in header]
    if missing:
        agg.die(tool, f"{path.name} sheet {SHEET!r} lacks {missing}; header is {header}")
    iz, iy, ip = header.index(ZIP_COL), header.index(YEAR_COL), header.index(PREMIUM_COL)

    years: dict[int, int] = {}
    premium: dict[str, float] = {}
    problems: list[str] = []
    blank = 0
    for n, row in enumerate(rows, start=2):
        if row is None or row[iz] is None:
            blank += 1
            continue
        z = agg.zip5(row[iz])
        if z is None:
            problems.append(f"row {n}: ZIP {row[iz]!r} is not a ZIP code")
            continue
        try:
            y = int(row[iy])
        except (TypeError, ValueError):
            problems.append(f"row {n}: year {row[iy]!r} is not a year")
            continue
        years[y] = years.get(y, 0) + 1
        if y != year:
            continue
        p = row[ip]
        if not isinstance(p, (int, float)) or isinstance(p, bool) or p <= 0:
            problems.append(f"row {n}: ZIP {z} {year} premium {p!r} is not a positive number")
            continue
        if z in premium:
            problems.append(f"row {n}: ZIP {z} appears twice for {year}")
            continue
        premium[z] = float(p)
    wb.close()
    agg.fail_if(problems[:25] + ([f"... and {len(problems) - 25} more"] if len(problems) > 25 else []), tool)
    if year not in years:
        agg.die(tool, f"{path.name} has no {year} rows; years present: {sorted(years)}")

    values = list(premium.values())
    stats = {
        "sheet": SHEET,
        "premium_field": PREMIUM_COL,
        "years_in_file": {str(y): c for y, c in sorted(years.items())},
        "zips_published_for_year": len(premium),
        "blank_rows_skipped": blank,
        "zip_premium_summary": {"min": round(min(values), 2), "median": round(statistics.median(values), 2),
                                "max": round(max(values), 2)},
        f"zips_above_{HIGH_PREMIUM}": sum(v > HIGH_PREMIUM for v in values),
    }
    return premium, stats


if __name__ == "__main__":
    sys.exit("fio_insurance.py is a module; run tax_insurance_to_mha.py")
