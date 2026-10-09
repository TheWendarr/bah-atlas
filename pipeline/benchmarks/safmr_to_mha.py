#!/usr/bin/env python3
"""
safmr_to_mha.py  (A5 -> B10 FMR at MHA)

Translate HUD's Small Area Fair Market Rents workbook (FY2026 revised) into
0-4 bedroom SAFMRs per Military Housing Area.

    python pipeline/benchmarks/safmr_to_mha.py \
        --xlsx data/raw/fy2026_safmrs_revised.xlsx \
        --crosswalk data/interim/bah_2026/zip_mha.csv \
        --out data/interim/benchmarks_2026

------------------------------------------------------------------------------
Input layout (HUD SAFMR workbook)
------------------------------------------------------------------------------
  One sheet, 18 columns: ZIP code, HUD area code, HUD area name, then for each
  of 0-4 bedrooms three columns: the SAFMR itself, its 90% payment standard, and
  its 110% payment standard. Header cells contain hard line breaks
  ("ZIP\\nCode", "SAFMR\\n2BR", "SAFMR\\n2BR - 90%\\nPayment\\nStandard"), so headers
  are matched after collapsing whitespace. Only the base SAFMR columns are used.
  A ZIP that straddles two HUD areas appears once per area.

------------------------------------------------------------------------------
Method
------------------------------------------------------------------------------
  * FMR is GROSS rent (rent plus tenant-paid utilities), so unlike ZORI it needs
    no utilities adjustment: it is the utilities-inclusive cross-check for B8.
  * A ZIP listed under more than one HUD area takes the mean of its rows (count
    recorded in the manifest).
  * MHA value per bedroom count = median across the MHA's ZIPs (see mha_agg.py).
    Each profile in config/profiles.json (A12) later selects its anchor_bedrooms
    row: 2BR for E-5 with dependents, 3BR for O-3 with dependents.
  * FMR is set at the 40th percentile of standard-quality recent-mover rents, by
    HUD's definition. It is a lower anchor than a typical market rent; say so
    wherever SAFMR and ZORI are shown side by side.

------------------------------------------------------------------------------
Output (written to --out)
------------------------------------------------------------------------------
  zip_safmr.csv            zip, mha, br0..br4, hud_rows        (ZIPs in real MHAs)
  mha_safmr.csv            mha, bedrooms, n_zips, n_zips_with_data, zip_coverage,
                           safmr_median, safmr_mean, safmr_min, safmr_max   (long: 5 rows per MHA)
  mha_safmr_manifest.json  source, sha256, sheet, header map, counts, warnings

------------------------------------------------------------------------------
Validation (loud)
------------------------------------------------------------------------------
  * a ZIP column and all five base SAFMR columns (0BR-4BR) are found; if not,
    the run stops and prints the headers it saw
  * ZIPs parse to 5 digits; values are numeric and within $100-$20,000
  * within each ZIP, SAFMRs do not fall as bedrooms rise (0BR <= 1BR <= ... 4BR);
    breaks are counted and reported, and stop the run past 1% of ZIPs
"""

from __future__ import annotations

import argparse
import re
import statistics
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import mha_agg as agg  # noqa: E402

TOOL = "safmr_to_mha"
SOURCE = ("U.S. Department of Housing and Urban Development (HUD) PD&R, "
          "Small Area Fair Market Rents FY2026 (revised, effective 21 May 2026)")
BEDROOMS = range(5)
BASE_COL = re.compile(r"^safmr\s*(\d)\s*br$")
ZIP_COL = re.compile(r"^zip(\s*code)?$")
LO, HI = 100.0, 20_000.0
MONOTONIC_TOLERANCE = 0.01


def norm(cell) -> str:
    return " ".join(str(cell or "").split()).lower()


def find_header(rows: list[tuple]) -> tuple[int, int, dict[int, int]] | None:
    """Return (header row index, zip column, {bedrooms: column}) or None."""
    for r_i, row in enumerate(rows):
        cells = [norm(c) for c in row]
        zip_cols = [i for i, c in enumerate(cells) if ZIP_COL.match(c)]
        beds = {}
        for i, c in enumerate(cells):
            m = BASE_COL.match(c)
            if m and int(m.group(1)) in BEDROOMS:
                beds.setdefault(int(m.group(1)), i)
        if zip_cols and len(beds) == len(BEDROOMS):
            return r_i, zip_cols[0], beds
    return None


def read_workbook(path: Path) -> tuple[dict[str, list[list[float]]], dict, list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        agg.die(TOOL, "openpyxl is required to read the HUD workbook: `pip install -r requirements.txt`")
    with path.open("rb") as fh:
        if fh.read(4) != b"PK\x03\x04":
            agg.die(TOOL, f"{path} is not an .xlsx workbook (re-run `make fetch`)")
    wb = load_workbook(path, read_only=True, data_only=True)
    problems: list[str] = []
    seen_headers: list[str] = []
    for ws in wb.worksheets:
        rows_iter = ws.iter_rows(values_only=True)
        head = []
        for _ in range(15):
            try:
                head.append(next(rows_iter))
            except StopIteration:
                break
        found = find_header(head)
        if not found:
            if head:
                seen_headers.append(f"[{ws.title}] " + " | ".join(norm(c) for c in head[0] if c is not None))
            continue
        h_i, zip_c, bed_c = found
        header_map = {"sheet": ws.title, "header_row": h_i + 1,
                      "zip": head[h_i][zip_c],
                      **{f"br{b}": head[h_i][c] for b, c in sorted(bed_c.items())}}
        by_zip: dict[str, list[list[float]]] = {}

        def take(row: tuple, line: int) -> None:
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                return
            z = agg.zip5(row[zip_c]) if zip_c < len(row) and row[zip_c] is not None else None
            if z is None:
                problems.append(f"row {line}: malformed ZIP {row[zip_c] if zip_c < len(row) else None!r}")
                return
            vals = []
            for b in BEDROOMS:
                c = bed_c[b]
                cell = row[c] if c < len(row) else None
                try:
                    v = float(str(cell).replace(",", "").replace("$", ""))
                except (TypeError, ValueError):
                    problems.append(f"row {line} ZIP {z}: {b}BR value {cell!r} is not a number")
                    return
                if not LO <= v <= HI:
                    problems.append(f"row {line} ZIP {z}: {b}BR ${v:,.0f} outside ${LO:,.0f}-${HI:,.0f}")
                    return
                vals.append(v)
            by_zip.setdefault(z, []).append(vals)

        for offset, row in enumerate(head[h_i + 1:], start=h_i + 2):
            take(row, offset)
        for offset, row in enumerate(rows_iter, start=len(head) + 1):
            take(row, offset)
        wb.close()
        return by_zip, header_map, problems
    wb.close()
    agg.die(TOOL, "no sheet has a ZIP column plus SAFMR 0BR-4BR columns. Headers seen:\n  "
            + "\n  ".join(seen_headers or ["(no rows)"]))
    raise AssertionError  # unreachable


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", type=Path, required=True, help="HUD SAFMR workbook")
    ap.add_argument("--crosswalk", type=Path, required=True, help="B1 zip_mha.csv")
    ap.add_argument("--out", type=Path, required=True, help="output folder")
    args = ap.parse_args()

    if not args.xlsx.exists():
        agg.die(TOOL, f"input not found: {args.xlsx} (run `make fetch`)")
    crosswalk, cw_stats = agg.load_crosswalk(args.crosswalk, TOOL)
    print(f"B10 SAFMR at MHA: reading {args.xlsx.name}")
    by_zip, header_map, problems = read_workbook(args.xlsx)
    agg.fail_if(problems[:25] + ([f"... and {len(problems) - 25} more"] if len(problems) > 25 else []), TOOL)
    if not by_zip:
        agg.die(TOOL, f"no data rows under the header in sheet {header_map['sheet']!r}")

    # Collapse ZIPs listed under more than one HUD area.
    zip_vals: dict[str, list[float]] = {}
    multi = 0
    for z, lists in by_zip.items():
        if len(lists) > 1:
            multi += 1
        zip_vals[z] = [statistics.fmean(v[b] for v in lists) for b in BEDROOMS]

    breaks = [z for z, v in zip_vals.items() if any(v[b] > v[b + 1] + 0.5 for b in range(4))]
    warnings: list[str] = []
    if breaks:
        msg = f"{len(breaks)} ZIP(s) have a SAFMR that falls as bedrooms rise (e.g. {', '.join(breaks[:5])})"
        if len(breaks) > MONOTONIC_TOLERANCE * len(zip_vals):
            agg.fail_if([msg + " -- more than 1% of ZIPs, so the columns are probably mislabelled"], TOOL)
        warnings.append(msg)
    if multi:
        warnings.append(f"{multi} ZIP(s) appear under more than one HUD area; their rows were averaged")

    out_rows: list[dict] = []
    per_bed_stats = {}
    for b in BEDROOMS:
        rows, stats = agg.aggregate({z: v[b] for z, v in zip_vals.items()}, crosswalk)
        per_bed_stats[f"br{b}"] = stats
        for r in rows:
            out_rows.append({
                "mha": r["mha"], "bedrooms": b, "n_zips": r["n_zips"],
                "n_zips_with_data": r["n_zips_with_data"], "zip_coverage": f"{r['zip_coverage']:.4f}",
                "safmr_median": agg.fmt(r["median"]), "safmr_mean": agg.fmt(r["mean"]),
                "safmr_min": agg.fmt(r["min"]), "safmr_max": agg.fmt(r["max"]),
            })
    out_rows.sort(key=lambda r: (r["mha"], r["bedrooms"]))
    agg.write_csv(args.out / "mha_safmr.csv",
                  ["mha", "bedrooms", "n_zips", "n_zips_with_data", "zip_coverage",
                   "safmr_median", "safmr_mean", "safmr_min", "safmr_max"], out_rows)

    zip_rows = [{"zip": z, "mha": crosswalk[z], **{f"br{b}": agg.fmt(v[b]) for b in BEDROOMS},
                 "hud_rows": len(by_zip[z])}
                for z, v in sorted(zip_vals.items()) if z in crosswalk]
    agg.write_csv(args.out / "zip_safmr.csv", ["zip", "mha", *[f"br{b}" for b in BEDROOMS], "hud_rows"], zip_rows)

    two = per_bed_stats["br2"]
    meds2 = [float(r["safmr_median"]) for r in out_rows if r["bedrooms"] == 2 and r["safmr_median"]]
    manifest = {
        "node": "B10",
        "source_node": "A5",
        "source": SOURCE,
        "unit": "USD per month (gross rent: rent plus tenant-paid utilities)",
        "input_file": args.xlsx.name,
        "input_sha256": agg.sha256(args.xlsx),
        "generated_at": agg.now_utc(),
        "header_map": {k: (" ".join(str(v).split()) if v is not None else None) for k, v in header_map.items()},
        "hud_zip_rows": sum(len(v) for v in by_zip.values()),
        "hud_unique_zips": len(zip_vals),
        "aggregation": "per bedroom count: median of ZIP SAFMRs across the MHA's ZIPs (unweighted); real MHAs only",
        "crosswalk": cw_stats,
        "coverage": two,
        "coverage_note": "Coverage is identical for every bedroom count because HUD publishes all five per ZIP; br2 shown.",
        "mha_2br_median_summary": ({"min": min(meds2), "median": statistics.median(meds2), "max": max(meds2)}
                                   if meds2 else None),
        "warnings": warnings,
        "outputs": ["zip_safmr.csv", "mha_safmr.csv"],
    }
    agg.write_manifest(args.out / "mha_safmr_manifest.json", manifest)

    print(f"  sheet {header_map['sheet']!r}, header row {header_map['header_row']}; "
          f"{manifest['hud_zip_rows']:,} rows, {len(zip_vals):,} unique ZIPs")
    print(f"  {two['mhas_with_data']} of {len(out_rows) // 5} MHAs have data; "
          f"median ZIP coverage {two['median_zip_coverage']:.0%}")
    if meds2:
        print(f"  MHA 2BR SAFMR: min {min(meds2):,.0f} / median {statistics.median(meds2):,.0f} / max {max(meds2):,.0f}")
    for w in warnings:
        print(f"  warning: {w}")
    print(f"  -> {args.out / 'mha_safmr.csv'}")


if __name__ == "__main__":
    main()
