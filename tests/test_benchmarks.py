"""
Tests for the 2.3 / 2.4 benchmark nodes (B8-B11) against small synthetic fixtures
that mirror the vendor layouts documented in docs/SOURCES.md.

    python -m unittest discover -s tests -v        (or: make test)

Each script runs as a subprocess, exactly as the Makefile calls it. Expected
numbers are worked out by hand in the comments so a reader can check them.
"""

from __future__ import annotations

import csv
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
BENCH = ROOT / "pipeline" / "benchmarks"


def run(script: str, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(BENCH / script), *args],
                          capture_output=True, text=True)


def read(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def write(path: Path, rows: list[list]) -> Path:
    with path.open("w", newline="", encoding="utf-8") as fh:
        csv.writer(fh, lineterminator="\n").writerows(rows)
    return path


class Fixture(unittest.TestCase):
    """Three real MHAs and one County Cost Group.

    AA001: ZIPs 00501, 00502       (both have Zillow data)
    AA002: ZIPs 10001, 10002, 10003 (10003 has only 2 months -> dropped at min 3)
    AA003: ZIP 20001               (no Zillow data at all)
    ZZ001: ZIP 30001               (CCG, excluded from aggregation)
    """

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)
        self.out = self.d / "out"
        self.crosswalk = write(self.d / "zip_mha.csv", [
            ["zip", "mha", "mha_class"],
            ["00501", "AA001", "mha"], ["00502", "AA001", "mha"],
            ["10001", "AA002", "mha"], ["10002", "AA002", "mha"], ["10003", "AA002", "mha"],
            ["20001", "AA003", "mha"],
            ["30001", "ZZ001", "ccg"],
        ])
        self.components = write(self.d / "rate_components.csv", [
            ["mha", "name", "rent_pct", "utilities_pct", "rent_share", "utilities_share"],
            ["AA001", "ONE", 80, 20, "0.8000", "0.2000"],
            ["AA002", "TWO", 75, 25, "0.7500", "0.2500"],
            ["AA003", "THREE", 85, 15, "0.8500", "0.1500"],
        ])
        head = ["RegionID", "SizeRank", "RegionName", "RegionType", "StateName", "State",
                "City", "Metro", "CountyName", "2025-11-30", "2025-12-31",
                "2026-01-31", "2026-02-28", "2026-03-31"]
        # Window 2026-01:2026-12 -> the three 2026 columns. ZIP 501 has lost its zeros.
        self.zori = write(self.d / "Zip_zori.csv", [
            head,
            [1, 1, "501", "zip", "", "NY", "", "", "", 900, 950, 1000, 1100, 1200],   # mean 1100
            [2, 2, "502", "zip", "", "NY", "", "", "", 0, 0, 2000, 2000, 2000],       # mean 2000 (the 2025 zeros fall outside the window)
            [3, 3, "10001", "zip", "", "NY", "", "", "", "", "", 3000, 3000, 3000],   # 3000
            [4, 4, "10002", "zip", "", "NY", "", "", "", "", "", 1500, 1600, 1700],   # 1600
            [5, 5, "10003", "zip", "", "NY", "", "", "", "", "", "", 9000, 9000],     # 2 months -> dropped
            [6, 6, "30001", "zip", "", "TX", "", "", "", "", "", 800, 800, 800],      # CCG
        ])

    def tearDown(self) -> None:
        self.tmp.cleanup()


class TestZillow(Fixture):
    def test_zori_b8(self) -> None:
        # The 2025 zeros in ZIP 502 sit outside the window, so the range check never sees them.
        p = run("zillow_to_mha.py", "--series", "zori", "--csv", str(self.zori),
                "--crosswalk", str(self.crosswalk), "--rate-components", str(self.components),
                "--out", str(self.out), "--window", "2026-01:2026-12")
        self.assertEqual(p.returncode, 0, p.stderr)
        rows = {r["mha"]: r for r in read(self.out / "mha_zori.csv")}
        self.assertEqual(sorted(rows), ["AA001", "AA002", "AA003"])           # CCG left out
        # AA001: median(1100, 2000) = 1550; / 0.80 = 1937.50; utilities 387.50
        self.assertEqual(rows["AA001"]["zori_median"], "1550.00")
        self.assertEqual(rows["AA001"]["zori_utilities_adjusted"], "1937.50")
        self.assertEqual(rows["AA001"]["utilities_usd"], "387.50")
        # AA002: 10003 dropped (2 months) -> median(3000, 1600) = 2300; coverage 2/3
        self.assertEqual(rows["AA002"]["zori_median"], "2300.00")
        self.assertEqual(rows["AA002"]["n_zips_with_data"], "2")
        self.assertEqual(rows["AA002"]["zip_coverage"], "0.6667")
        # AA003 has no data but still gets a row
        self.assertEqual(rows["AA003"]["n_zips_with_data"], "0")
        self.assertEqual(rows["AA003"]["zori_median"], "")
        man = json.loads((self.out / "mha_zori_manifest.json").read_text())
        self.assertEqual(man["window_months_used"], ["2026-01", "2026-02", "2026-03"])
        self.assertEqual(man["coverage"]["mhas_without_data"], ["AA003"])
        self.assertEqual(man["crosswalk"]["zips_skipped_by_class"], {"ccg": 1})
        zips = {r["zip"] for r in read(self.out / "zip_zori.csv")}
        self.assertIn("00501", zips)                                          # zero-padded
        self.assertNotIn("30001", zips)

    def test_zori_requires_components(self) -> None:
        p = run("zillow_to_mha.py", "--series", "zori", "--csv", str(self.zori),
                "--crosswalk", str(self.crosswalk), "--out", str(self.out))
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("--rate-components", p.stderr)

    def test_zhvi_b9_and_short_window(self) -> None:
        with self.zori.open(newline="") as fh:
            rows = list(csv.reader(fh))
        for r in rows[1:]:                      # rents x 200 -> home values
            r[9:] = [str(int(v) * 200) if v else "" for v in r[9:]]
        zhvi = write(self.d / "Zip_zhvi.csv", rows)
        p = run("zillow_to_mha.py", "--series", "zhvi", "--csv", str(zhvi),
                "--crosswalk", str(self.crosswalk), "--out", str(self.out), "--window", "2026-01:2026-12")
        self.assertEqual(p.returncode, 0, p.stderr)
        rows = {r["mha"]: r for r in read(self.out / "mha_zhvi.csv")}
        self.assertEqual(rows["AA001"]["zhvi_median"], "310000.00")       # median(220k, 400k)
        p = run("zillow_to_mha.py", "--series", "zhvi", "--csv", str(zhvi),
                "--crosswalk", str(self.crosswalk), "--out", str(self.out), "--window", "2026-03:2026-12")
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("need at least 3", p.stderr)

    def test_out_of_range_value_fails(self) -> None:
        bad = self.d / "bad.csv"
        with self.zori.open(newline="") as fh:
            rows = list(csv.reader(fh))
        rows[1][-1] = "5"                       # a $5 rent
        write(bad, rows)
        p = run("zillow_to_mha.py", "--series", "zori", "--csv", str(bad),
                "--crosswalk", str(self.crosswalk), "--rate-components", str(self.components),
                "--out", str(self.out), "--window", "2026-01:2026-12")
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("outside plausible range", p.stderr)

    def test_seasonal_luxury_rent_warns(self) -> None:
        # A Hamptons-style $95k summer rent is real: kept, listed, and does not stop the run.
        hot = self.d / "hot.csv"
        with self.zori.open(newline="") as fh:
            rows = list(csv.reader(fh))
        rows[4][-1] = "95000"                   # ZIP 10002 (AA002), 2026-03
        write(hot, rows)
        p = run("zillow_to_mha.py", "--series", "zori", "--csv", str(hot),
                "--crosswalk", str(self.crosswalk), "--rate-components", str(self.components),
                "--out", str(self.out), "--window", "2026-01:2026-12")
        self.assertEqual(p.returncode, 0, p.stderr)
        man = json.loads((self.out / "mha_zori_manifest.json").read_text())
        self.assertEqual(man["extreme_zips"], {"10002": 95000.0})
        self.assertTrue(any("10002 (AA002, peak 95,000)" in w for w in man["warnings"]))


class TestSafmr(Fixture):
    def make_xlsx(self, headers: list[str], rows: list[list]) -> Path:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(r)
        path = self.d / "safmr.xlsx"
        wb.save(path)
        return path

    HEAD = ["ZIP\nCode", "HUD Area Code", "HUD Fair Market Rent Area Name"] + [
        h for b in range(5) for h in (f"SAFMR\n{b}BR", f"SAFMR\n{b}BR - 90%\nPayment\nStandard",
                                      f"SAFMR\n{b}BR - 110%\nPayment\nStandard")]

    @staticmethod
    def row(z, area, base: list[int]) -> list:
        return [z, area, "Area"] + [v for b in base for v in (b, round(b * 0.9), round(b * 1.1))]

    def test_b10(self) -> None:
        x = self.make_xlsx(self.HEAD, [
            self.row(501, "A1", [1000, 1100, 1300, 1700, 2000]),      # integer ZIP, lost zeros
            self.row("00502", "A1", [1200, 1300, 1500, 1900, 2200]),
            self.row("10001", "A2", [2000, 2100, 2400, 3000, 3300]),
            self.row("10001", "A3", [2200, 2300, 2600, 3200, 3500]),  # straddles two HUD areas
        ])
        p = run("safmr_to_mha.py", "--xlsx", str(x), "--crosswalk", str(self.crosswalk), "--out", str(self.out))
        self.assertEqual(p.returncode, 0, p.stderr)
        rows = {(r["mha"], r["bedrooms"]): r for r in read(self.out / "mha_safmr.csv")}
        self.assertEqual(len(rows), 15)                                   # 3 MHAs x 5 bedroom counts
        self.assertEqual(rows[("AA001", "2")]["safmr_median"], "1400.00")  # median(1300, 1500)
        self.assertEqual(rows[("AA002", "2")]["safmr_median"], "2500.00")  # mean of two area rows
        self.assertEqual(rows[("AA002", "2")]["zip_coverage"], "0.3333")
        self.assertEqual(rows[("AA003", "3")]["safmr_median"], "")
        man = json.loads((self.out / "mha_safmr_manifest.json").read_text())
        self.assertEqual(man["header_map"]["br2"], "SAFMR 2BR")             # base column, not payment std
        self.assertTrue(any("more than one HUD area" in w for w in man["warnings"]))

    def test_missing_columns_print_headers(self) -> None:
        x = self.make_xlsx(["ZIP Code", "Two Bedroom"], [["00501", 1000]])
        p = run("safmr_to_mha.py", "--xlsx", str(x), "--crosswalk", str(self.crosswalk), "--out", str(self.out))
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("two bedroom", p.stderr)                            # shows what it saw

    def test_mislabelled_bedrooms_fail(self) -> None:
        x = self.make_xlsx(self.HEAD, [self.row("00501", "A1", [2000, 1800, 1600, 1400, 1200])])
        p = run("safmr_to_mha.py", "--xlsx", str(x), "--crosswalk", str(self.crosswalk), "--out", str(self.out))
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("mislabelled", p.stderr)


class TestOwnership(Fixture):
    def test_b11(self) -> None:
        zhvi = write(self.d / "mha_zhvi.csv", [
            ["mha", "n_zips", "n_zips_with_data", "zip_coverage", "zhvi_median"],
            ["AA001", 2, 2, "1.0000", "300000.00"], ["AA002", 3, 0, "0.0000", ""],
            ["AA003", 1, 1, "1.0000", "200000.00"]])
        zori = write(self.d / "mha_zori.csv", [
            ["mha", "utilities_usd"], ["AA001", "387.50"], ["AA002", "400.00"], ["AA003", ""]])
        # Blank cells hold a single space, as in Freddie Mac's file; one row outside the window.
        pmms = write(self.d / "PMMS_history.csv", [
            ["date", "pmms30", "pmms30p", "pmms15"],
            ["12/26/2025", "9.00", " ", " "],
            ["1/1/2026", "6.40", " ", " "], ["1/8/2026", "6.60", " ", " "],
            ["2/5/2026", "6.40", " ", " "], ["3/5/2026", "6.60", " ", " "]])
        # The national 2025 placeholder values, now per MHA, keep the hand-worked numbers below.
        taxins = write(self.d / "mha_tax_insurance.csv", [
            ["mha", "effective_tax_rate", "tax_source", "insurance_annual_usd", "insurance_source"],
            ["AA001", "0.008900", "locality", "1569.00", "locality"],
            ["AA002", "0.012000", "locality", "2000.00", "state_fallback"],
            ["AA003", "0.010000", "state_fallback", "1800.00", "locality"]])
        p = run("ownership_cost.py", "--zhvi", str(zhvi), "--zori", str(zori), "--pmms", str(pmms),
                "--tax-insurance", str(taxins), "--out", str(self.out), "--window", "2026-01:2026-12")
        self.assertEqual(p.returncode, 0, p.stderr)
        rows = {r["mha"]: r for r in read(self.out / "mha_ownership.csv")}
        a = rows["AA001"]
        self.assertEqual(a["rate_pct"], "6.500")                            # 9.00 week excluded
        self.assertEqual(a["loan_amount"], "306450.00")                     # 300k + 2.15% fee, 0% down
        # $300,000 at 6.5%/30y is $1,896.20/mo, so $306,450 is $1,936.97
        self.assertAlmostEqual(float(a["pi_monthly"]), 1936.97, delta=0.02)
        self.assertEqual(a["tax_monthly"], "222.50")                        # 300k x 0.89% / 12
        self.assertEqual(a["insurance_monthly"], "130.75")                  # 1569 / 12
        self.assertEqual((a["tax_source"], a["insurance_source"]), ("locality", "locality"))
        self.assertEqual(rows["AA003"]["tax_monthly"], "166.67")            # 200k x 1.0% / 12
        self.assertEqual(rows["AA003"]["tax_source"], "state_fallback")
        self.assertAlmostEqual(float(a["own_housing_monthly"]), 1936.97 + 222.50 + 130.75, delta=0.03)
        self.assertAlmostEqual(float(a["own_total_monthly"]), float(a["own_housing_monthly"]) + 387.50, delta=0.01)
        self.assertEqual(rows["AA002"].get("own_housing_monthly", ""), "")  # no home value
        self.assertEqual(rows["AA003"]["own_total_monthly"], "")            # no utilities estimate
        man = json.loads((self.out / "mha_ownership_manifest.json").read_text())
        self.assertEqual(man["rate"]["weeks"], 4)
        self.assertFalse(any("placeholder" in w for w in man["warnings"]))
        self.assertEqual(man["input_sha256"]["A8"], __import__("hashlib").sha256(taxins.read_bytes()).hexdigest())
        self.assertEqual(man["tax_insurance_sources"]["tax_source"], {"locality": 1, "state_fallback": 1})
        self.assertTrue(any("not scaled" in x for x in man["limitations"]))

    def test_bad_config_fails(self) -> None:
        cfg = json.loads((ROOT / "config" / "ownership.json").read_text())
        cfg["loan"]["funding_fee_share"] = 2.15                            # percent typed as share
        cfg["insurance"]["scaled_to_bah_year"] = True                       # decision N2 says as published
        bad = self.d / "ownership.json"
        bad.write_text(json.dumps(cfg))
        p = run("ownership_cost.py", "--zhvi", str(self.d / "x.csv"), "--pmms", str(self.d / "y.csv"),
                "--tax-insurance", str(self.d / "z.csv"), "--out", str(self.out), "--config", str(bad))
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("loan.funding_fee_share", p.stderr)
        self.assertIn("scaled_to_bah_year", p.stderr)

    def test_bad_tax_rate_fails(self) -> None:
        zhvi = write(self.d / "mha_zhvi.csv", [["mha", "zhvi_median"], ["AA001", "300000.00"]])
        taxins = write(self.d / "mha_tax_insurance.csv", [
            ["mha", "effective_tax_rate", "tax_source", "insurance_annual_usd", "insurance_source"],
            ["AA001", "0.89", "locality", "1569.00", "locality"]])        # percent typed as share
        p = run("ownership_cost.py", "--zhvi", str(zhvi), "--pmms", str(self.d / "y.csv"),
                "--tax-insurance", str(taxins), "--out", str(self.out))
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("effective_tax_rate", p.stderr)


class TestTaxInsurance(Fixture):
    """A8 rollup. ACS and FIO values per ZIP; every real MHA is in state "AA".

    ACS (taxes, value): 00501 (3000, 300000), 00502 (1000, 200000)  -> AA001 4000/500000 = 0.008
                        10001 (6000, 400000), 10002 tax is an annotation code (-666666666),
                        10003 value 0                                -> AA002 6000/400000 = 0.015
                        20001 absent                                 -> AA003 state (10000/900000 = 0.011111)
                        30001 (CCG)                                  -> ignored
    FIO 2022 premiums:  00501 1200, 00502 1400                       -> AA001 median 1300
                        10001 2000, 10002 2600, 10003 3000           -> AA002 median 2600
                        20001 absent                                 -> AA003 state median of the five = 2000
    """

    def make_acs(self, header: list[str] | None = None, vintage: int = 2024) -> Path:
        raw = self.d / "raw"
        raw.mkdir(exist_ok=True)
        header = header or ["B25090_001E", "B25079_001E", "B01003_001E", "zip code tabulation area"]
        table = [header,
                 ["3000", "300000", "900", "00501"], ["1000", "200000", "800", "00502"],
                 ["6000", "400000", "1500", "10001"], ["-666666666", "250000", "400", "10002"],
                 ["0", "0", "10", "10003"], ["9000", "500000", "2000", "30001"]]
        acs = raw / "acs5_2024_zcta.json"
        acs.write_text(json.dumps(table))
        (raw / "fetch_manifest.json").write_text(json.dumps({"A9": {
            "filename": acs.name, "source_url": f"https://api.census.gov/data/{vintage}/acs/acs5?get=x",
            "sha256": "0" * 64}}))
        return acs

    def make_fio(self, year: int = 2022) -> Path:
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.title = "Disclaimer & Background"
        ws = wb.create_sheet("Supporting Underlying Metrics")
        ws.append(["ZIP Code", "Year", "Policy Decile Grouping", "Claim Frequency", "Claim Severity",
                   "Loss Ratio", "Premiums Per Policy", "Nonrenewal Rate", "Nonpayment Cancellation Rate",
                   "Other than Nonpayment Cancellation Rate"])
        for z, prem in ((501, 1200), (502, 1400), (10001, 2000), (10002, 2600), (10003, 3000)):
            ws.append([z, year, 5, 0.03, 10000, 0.4, prem, 0.01, 0.02, 0.01])
        ws.append([501, year - 1, 5, 0.03, 10000, 0.4, 999, 0.01, 0.02, 0.01])   # other year: ignored
        ws.append([None] * 10)                                                    # trailing blank row
        path = self.d / "fio.xlsx"
        wb.save(path)
        return path

    def run_taxins(self, acs: Path, fio: Path) -> subprocess.CompletedProcess:
        return run("tax_insurance_to_mha.py", "--acs", str(acs), "--fio", str(fio),
                   "--crosswalk", str(self.crosswalk), "--out", str(self.out))

    def test_rates_premiums_and_fallbacks(self) -> None:
        p = self.run_taxins(self.make_acs(), self.make_fio())
        self.assertEqual(p.returncode, 0, p.stderr)
        rows = {r["mha"]: r for r in read(self.out / "mha_tax_insurance.csv")}
        self.assertEqual(sorted(rows), ["AA001", "AA002", "AA003"])          # CCG excluded
        a, b, c = rows["AA001"], rows["AA002"], rows["AA003"]
        self.assertEqual((a["effective_tax_rate"], a["tax_zips_with_data"], a["tax_source"]),
                         ("0.008000", "2", "locality"))
        self.assertEqual((b["effective_tax_rate"], b["tax_zips_with_data"]), ("0.015000", "1"))
        self.assertEqual((c["effective_tax_rate"], c["tax_source"]), ("0.011111", "state_fallback"))
        self.assertEqual((a["insurance_annual_usd"], a["insurance_source"]), ("1300.00", "locality"))
        self.assertEqual(b["insurance_annual_usd"], "2600.00")
        self.assertEqual((c["insurance_annual_usd"], c["insurance_source"]), ("2000.00", "state_fallback"))
        man = json.loads((self.out / "mha_tax_insurance_manifest.json").read_text())
        self.assertEqual(man["nodes"], ["A8"])
        self.assertEqual(sorted(man["inputs"]), ["A8", "A9", "B1"])
        pt, ins = man["property_tax"], man["insurance"]
        self.assertEqual((pt["acs_vintage"], pt["acs_period"]), (2024, "2020-2024"))
        self.assertEqual((pt["zctas_in_file"], pt["zctas_usable"], pt["zctas_missing_estimate"],
                          pt["zctas_zero_value"]), (6, 4, 1, 1))
        self.assertEqual((pt["mhas_state_fallback"], ins["mhas_state_fallback"]), (["AA003"], ["AA003"]))
        self.assertEqual((ins["zips_published_for_year"], ins["zips_in_real_mhas"]), (5, 5))
        self.assertEqual(man["state_values_used"]["AA"], {"tax_rate": 0.011111, "premium": 2000.0})
        self.assertEqual(len(man["warnings"]), 2)                           # one per fallback kind

    def test_vintage_mismatch_stops(self) -> None:
        p = self.run_taxins(self.make_acs(vintage=2023), self.make_fio())
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("2023", p.stderr)

    def test_acs_without_tax_variable_stops(self) -> None:
        acs = self.make_acs(header=["B25079_001E", "B25090_00XE", "B01003_001E", "zip code tabulation area"])
        p = self.run_taxins(acs, self.make_fio())
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("B25090_001E", p.stderr)

    def test_fio_without_the_year_stops(self) -> None:
        p = self.run_taxins(self.make_acs(), self.make_fio(year=2021))
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("no 2022 rows", p.stderr)

    def test_state_without_any_data_stops(self) -> None:
        write(self.crosswalk, [["zip", "mha", "mha_class"], ["00501", "AA001", "mha"],
                               ["40001", "BB001", "mha"]])
        p = self.run_taxins(self.make_acs(), self.make_fio())
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("BB001: no ACS tax data in the MHA or its state BB", p.stderr)


class TestDensity(Fixture):
    """B24 population density. ACS population (TestTaxInsurance.make_acs) and ZCTA land area:

        ZCTA   pop    land km2   density
        00501   900      2         450      AA001: 1700 / 10 = 170.000
        00502   800      8         100             pw (900*450 + 800*100) / 1700 = 285.294
        10001  1500      1        1500      AA002: 1900 / 5 = 380.000
        10002   400      4         100             pw (1500*1500 + 400*100) / 1900 = 1205.263
        10003    10      0          -       (zero land: skipped)
        20001     -      5          -       AA003: no population -> blank, warned
    """

    def make_zcta(self, land_column: str = "ALAND20") -> Path:
        import geopandas as gpd
        from shapely.geometry import box
        land = {"00501": 2e6, "00502": 8e6, "10001": 1e6, "10002": 4e6, "10003": 0, "20001": 5e6, "30001": 9e6}
        gdf = gpd.GeoDataFrame({"ZCTA5CE20": list(land), land_column: [int(v) for v in land.values()]},
                               geometry=[box(i, 0, i + 1, 1) for i in range(len(land))], crs="EPSG:4326")
        path = self.d / "zcta.geojson"
        gdf.to_file(path, driver="GeoJSON")
        return path

    def run_density(self, acs: Path, zcta: Path) -> subprocess.CompletedProcess:
        return run("density_to_mha.py", "--acs", str(acs), "--zcta", str(zcta),
                   "--crosswalk", str(self.crosswalk), "--out", str(self.out))

    def test_b24(self) -> None:
        p = self.run_density(TestTaxInsurance.make_acs(self), self.make_zcta())
        self.assertEqual(p.returncode, 0, p.stderr)
        rows = {r["mha"]: r for r in read(self.out / "mha_density.csv")}
        self.assertEqual(sorted(rows), ["AA001", "AA002", "AA003"])          # CCG excluded
        a, b, c = rows["AA001"], rows["AA002"], rows["AA003"]
        self.assertEqual((a["n_zctas_with_data"], a["population"], a["land_km2"], a["density_per_km2"],
                          a["pw_density_per_km2"]), ("2", "1700", "10.000", "170.000", "285.294"))
        self.assertEqual((b["n_zctas_with_data"], b["density_per_km2"], b["pw_density_per_km2"]),
                         ("2", "380.000", "1205.263"))
        self.assertEqual((c["n_zctas_with_data"], c["density_per_km2"]), ("0", ""))
        man = json.loads((self.out / "mha_density_manifest.json").read_text())
        self.assertEqual(man["nodes"], ["B24"])
        self.assertEqual(sorted(man["inputs"]), ["A2", "A9", "B1"])
        self.assertEqual((man["acs_vintage"], man["zctas_zero_land"], man["mhas_without_data"]),
                         (2024, 1, ["AA003"]))
        self.assertIn("mha_density.csv", man["output_sha256"])
        self.assertEqual(len(man["warnings"]), 1)

    def test_zcta_file_without_land_area_stops(self) -> None:
        p = self.run_density(TestTaxInsurance.make_acs(self), self.make_zcta(land_column="AREA"))
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("ALAND20", p.stderr)

    def test_acs_without_population_stops(self) -> None:
        acs = TestTaxInsurance.make_acs(self, header=["B25090_001E", "B25079_001E", "B01003_00XE",
                                                      "zip code tabulation area"])
        p = self.run_density(acs, self.make_zcta())
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("B01003_001E", p.stderr)


if __name__ == "__main__":
    unittest.main()
